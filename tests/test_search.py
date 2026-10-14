import datetime
import json
import os
import sqlite3
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from local_register_search.database import connect, needs_search_update
from local_register_search.indexer import sync
from local_register_search.search import search
from local_register_search.search_terms import parse_query, match_spans
from local_register_search.desktop_features import export_records


class SearchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db_path = self.root / 'index.sqlite3'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Register'
        sheet.append(['Name', 'Amount', 'Phone', 'Date', 'Notes'])
        sheet.append(['Rishikesh', 2973.6, '98765-43210', datetime.date(2026, 10, 13), '100% A_B'])
        sheet.append(['Adit 1', 90.3, '+91 98765 43210', datetime.date(2025, 1, 2), 'prefix needle suffix'])
        sheet.append(['Adit 10', 80.3, '9876543210', datetime.date(2024, 2, 29), '1000 AXB'])
        for i in range(105):
            sheet.append([f'Person {i}', i, '', '', 'bulk record'])
        sheet = book.create_sheet('Other')
        sheet.append(['Name', 'Amount'])
        sheet.append(['Other', 12])
        book.save(self.root / 'register.xlsx')
        book.close()
        result = sync(self.db_path, [str(self.root)], threading.Event())
        self.assertEqual(result['indexed'], 1, result)

    def test_queries_and_modes(self):
        self.assertEqual(parse_query('"Adit 1" 90.3'), ['adit 1', '90.3'])
        self.assertEqual(parse_query('13 Oct 2026'), ['13 oct 2026'])
        self.assertEqual(parse_query('12 34'), ['12', '34'])
        self.assertEqual(search(self.db_path, 'rishikesh 2973.6')[0], 1)
        self.assertEqual(search(self.db_path, '"Adit 1" 90.3')[0], 1)
        self.assertEqual(search(self.db_path, 'rishikesh 90.3', 'Any word')[0], 2)
        self.assertEqual(search(self.db_path, 'rishikesh 90.3')[0], 0)
        self.assertEqual(search(self.db_path, 'ishik')[0], 1)
        self.assertEqual(search(self.db_path, '100%')[0], 1)
        self.assertEqual(search(self.db_path, 'A_B')[0], 1)
        for query in ('', ' ', 'a', '""', '-', '/', ',', '|'):
            self.assertEqual(search(self.db_path, query), (0, []))

    def test_phones_dates_headers_filters(self):
        for phone in ('9876543210', '98765-43210', '+91 98765 43210', '(98765) 43210'):
            self.assertEqual(search(self.db_path, phone)[0], 3, phone)
        for date in ('13-10-2026', '13/10/2026', '2026-10-13', '13.10.2026', '13 Oct 2026', '13 October 2026'):
            self.assertEqual(search(self.db_path, date)[0], 1, date)
        self.assertEqual(search(self.db_path, 'Name')[0], 0)
        row = search(self.db_path, 'rishikesh')[1][0]
        self.assertEqual(row['rownum'], 2)
        self.assertEqual(json.loads(row['headers'])['A'], 'Name')
        self.assertEqual(search(self.db_path, 'other', sheet='Register')[0], 0)
        self.assertEqual(search(self.db_path, 'other', sheet='Other')[0], 1)

    def test_pagination_export_and_sync(self):
        pages = [search(self.db_path, 'bulk', page=i) for i in range(3)]
        self.assertEqual([len(rows) for total, rows in pages], [50, 50, 5])
        self.assertTrue(all(total == 105 for total, rows in pages))
        ids = [row['id'] for _, rows in pages for row in rows]
        self.assertEqual(len(set(ids)), 105)
        matching = search(self.db_path, 'bulk', page_size=-1)[1]
        self.assertEqual(export_records(self.root / 'results.csv', matching), 105)
        self.assertEqual(sync(self.db_path, [str(self.root)], threading.Event())['skipped'], 1)

    def test_migration_forces_unchanged_file_resync(self):
        db = connect(self.db_path)
        with db:
            db.execute("UPDATE rows SET row_text='',row_digits=''")
        self.assertTrue(needs_search_update(db))
        db.close()
        self.assertEqual(sync(self.db_path, [str(self.root)], threading.Event())['indexed'], 1)
        db = connect(self.db_path)
        self.assertFalse(needs_search_update(db))
        db.close()
        legacy = self.root / 'legacy.sqlite3'
        db = sqlite3.connect(legacy)
        db.execute('CREATE TABLE rows(id INTEGER PRIMARY KEY,path TEXT,sheet TEXT,rownum INTEGER,hidden INTEGER,cells TEXT,text TEXT,normalized TEXT)')
        db.execute("INSERT INTO rows VALUES(1,'old','Sheet',2,0,'{}','old','old')")
        db.commit()
        db.close()
        db = connect(legacy)
        self.assertTrue(needs_search_update(db))
        self.assertTrue({'row_text', 'row_digits'} <= {row['name'] for row in db.execute('PRAGMA table_info(rows)')})
        db.close()

    def test_highlight_spans(self):
        for value, query in [('Rishikesh 2973.6', 'rishikesh 2973.6'), ('Adit 1 / 90.3', '"Adit 1" 90.3')]:
            self.assertEqual(len(match_spans(value, query)), 2)
        self.assertEqual(match_spans('98765-43210', '9876543210'), [(0, 11)])
        self.assertEqual(match_spans('9876543210', '+91 98765 43210'), [(0, 10)])
        self.assertEqual(match_spans('13 Oct 2026', '13/10/2026'), [(0, 11)])
        self.assertEqual(match_spans('Straße', 'STRASSE'), [(0, 6)])

    def test_search_requires_synced_backend_records(self):
        from PySide6.QtWidgets import QApplication
        from local_register_search.ui import Window
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        app = QApplication.instance() or QApplication([])
        profile = self.root / 'empty-profile'
        profile.mkdir()
        with patch('local_register_search.ui.data_dir', return_value=profile):
            window = Window()
        self.addCleanup(window.close)
        self.assertEqual(window.empty_title.text(), 'No searchable data')
        self.assertIn('add a folder', window.empty_message.text())
        self.assertIn('color:#b42318', window.empty_message.styleSheet())
        self.assertEqual(window.result_count.text(), '')
        self.assertEqual(window.count.text(), '')
        window.set_dark_mode(True)
        self.assertIn('color:#ffb4ab', window.empty_message.styleSheet())
        window.set_dark_mode(False)
        window.query.setText('rishikesh')
        window.indexed = 100  # A stale UI count must not permit an empty backend search.
        with patch.object(window, 'launch') as launch:
            window.submit_search()
            launch.assert_not_called()
        self.assertEqual(window.table.rowCount(), 0)
        self.assertNotIn('No matching', window.result_count.text())
        self.assertEqual(window.result_count.text(), '')
        self.assertEqual(window.count.text(), '')
        self.assertFalse(window.export_button.isEnabled())
        self.assertTrue(all(not button.isEnabled() for button in window.page_buttons))
        window.folders = [str(self.root)]
        window.update_empty_state()
        self.assertIn('Click Sync files', window.empty_message.text())
        window.scanning = True
        window.update_empty_state()
        self.assertIn('Files are syncing', window.empty_message.text())
        window.scanning = False
        sync(window.path, [str(self.root)], threading.Event())
        window.refresh_filters()
        with patch.object(window, 'launch') as launch:
            window.submit_search()
            launch.assert_called_once()
        self.assertNotIn('color:#b42318', window.result_count.styleSheet())
        window.search_busy = False
        window.search_done(window.generation, (0, []))
        self.assertIn('No matching rows', window.result_count.text())

    def test_window_empty_state_paging_sync_and_migration(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtWidgets import QApplication
        from PySide6.QtTest import QTest
        from local_register_search.ui import Window
        app = QApplication.instance() or QApplication([])
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        window.path = self.db_path
        window.refresh_filters()
        self.assertEqual(window.table.rowCount(), 0)
        self.assertEqual(window.empty_message.text(), 'Type a keyword and press Enter.')

        self.assertEqual(window.page_size.count(), 1)
        self.assertEqual(window.page_size.currentData(), 50)

        def settle():
            for _ in range(400):
                app.processEvents()
                if not window.search_busy and not window.timer.isActive() and not window.jobs:
                    return
                QTest.qWait(10)
            self.fail('Search worker did not finish')

        window.query.setText('bulk')
        window.submit_search()
        settle()
        self.assertEqual(window.table.rowCount(), 50)
        self.assertTrue(window.result_count.text().startswith('Showing 1-50 of 105 results'))
        window.change_page(1)
        settle()
        self.assertTrue(window.result_count.text().startswith('Showing 51-100 of 105 results'))
        window.change_page(1)
        settle()
        self.assertEqual(window.table.rowCount(), 5)
        self.assertFalse(window.page_buttons[1].isEnabled())
        window.query.setText('a')
        self.assertEqual(window.table.rowCount(), 0)
        self.assertFalse(window.export_button.isEnabled())
        window.query.setText('rishikesh')
        window.submit_search()
        settle()
        window.sync_done(dict(indexed=1, skipped=0, removed=0, errors=[], cancelled=False))
        settle()
        self.assertEqual(window.table.rowCount(), 0)
        self.assertEqual(window.query.text(), '')
        db = connect(self.db_path)
        with db:
            db.execute("UPDATE rows SET row_text=''")
        db.close()
        window.refresh_filters()
        self.assertEqual(window.status.text(), 'Search index update ho raha hai. Sync files dobara dabayein')
        self.assertEqual(window.empty_message.text(), 'Type a keyword and press Enter.')


    def test_searchable_dropdowns_scope_sort_and_geometry(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QGuiApplication
        from PySide6.QtWidgets import QApplication
        from local_register_search.ui import Window
        from local_register_search.filter_dropdown import natural_key
        book = Workbook()
        book.remove(book.active)
        names = [f'TP-{i}' for i in range(43, 8, -1)] + ['Sheet10', 'Sheet2', 'Sheet1', 'A very long worksheet name']
        for name in names:
            sheet = book.create_sheet(name)
            sheet.append(['Name', 'Value'])
            sheet.append(['Example', 1])
        source = self.root / 'many-sheets.xlsx'
        book.save(source)
        book.close()
        sync(self.db_path, [str(self.root)], threading.Event())
        app = QApplication.instance() or QApplication([])
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        window.show()
        app.processEvents()
        sheets = window.sheet_filter
        self.assertEqual(sheets.itemText(0), 'All worksheets')
        actual = [sheets.itemText(i) for i in range(1, sheets.count())]
        self.assertEqual(actual, sorted(set(actual), key=natural_key))
        self.assertLess(actual.index('TP-9'), actual.index('TP-10'))
        self.assertLess(actual.index('Sheet2'), actual.index('Sheet10'))
        window.file_filter.setCurrentIndex(window.file_filter.findData(str(source.resolve())))
        window.timer.stop()
        self.assertEqual({sheets.itemText(i) for i in range(1, sheets.count())}, set(names))
        sheets.showPopup()
        app.processEvents()
        self.addCleanup(sheets.hidePopup)
        self.assertTrue(sheets.list_view.verticalScrollBar().maximum() > 0)
        self.assertLessEqual(sheets.list_view.viewport().height(), 12 * max(32, sheets.fontMetrics().height() + 16))
        screen = QGuiApplication.screenAt(sheets.popup.pos()) or QGuiApplication.primaryScreen()
        self.assertTrue(screen.availableGeometry().contains(sheets.popup.geometry()))
        sheets.filter_input.setText('tp-3')
        filtered = [sheets.list_view.item(i).text() for i in range(sheets.list_view.count())]
        self.assertEqual(filtered[0], 'All worksheets')
        self.assertTrue(all('tp-3' in name.casefold() for name in filtered[1:]))
        sheets.filter_input.setText('does-not-exist')
        self.assertEqual(sheets.list_view.item(1).text(), 'No sheets found')
        self.assertFalse(sheets.list_view.item(1).flags() & Qt.ItemFlag.ItemIsEnabled)
        sheets.hidePopup()
        window.file_filter.setCurrentIndex(0)
        window.timer.stop()
        self.assertGreater(sheets.count(), len(names) + 1)
        for combo, message in [(window.file_filter, 'No files found'), (window.folder_filter, 'No folders found')]:
            combo.showPopup()
            combo.filter_input.setText('not-found-xyz')
            self.assertEqual(combo.list_view.item(1).text(), message)
            combo.hidePopup()

    def test_section_context_exact_rows_and_mixed_headers(self):
        import hashlib
        from openpyxl.styles import Font, PatternFill
        from local_register_search.chip_input import quoted
        book = Workbook()
        sheet = book.active
        sheet.title = 'Sections'
        sheet.append(['Field', 'Value', 'Notes'])
        sheet.cell(18, 1, 'B - Client Data').font = Font(bold=True)
        for row, label in [(19, 'Given Name'), (20, 'Company'), (21, 'Function'), (22, 'Phone')]:
            sheet.cell(row, 1, label).font = Font(bold=True)
            sheet.cell(row, 2, 'Client detail')
        sheet.merge_cells('A29:C29')
        sheet.cell(29, 1, 'C - Archive Responsible')
        sheet.cell(29, 1).fill = PatternFill('solid', fgColor='FFFF00')
        for row, label in [(30, 'Given Name'), (31, 'Company'), (32, 'Function'), (33, 'Phone')]:
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, 'Archive detail')
        other = book.create_sheet('Plain')
        other.append(['Different field', 'Different value'])
        other.append(['Company', 'Other detail'])
        for row in range(3, 6):
            other.append(['Function', 'Plain detail'])
        source = self.root / 'sections.xlsx'
        book.save(source)
        book.close()
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        result = sync(self.db_path, [str(self.root)], threading.Event())
        self.assertEqual(result['indexed'], 1, result)
        for section, expected_row in [('B - Client Data', 20), ('C - Archive Responsible', 31)]:
            query = quoted('Company') + ' ' + quoted(section)
            total, records = search(self.db_path, query)
            self.assertEqual(total, 1)
            self.assertEqual(records[0]['rownum'], expected_row)
            self.assertEqual(records[0]['row_section'], section)
            self.assertEqual(search(self.db_path, query, include_section=False)[0], 0)
        self.assertEqual(search(self.db_path, '"Company"', sheet='Plain')[1][0]['row_section'], '')
        db = connect(self.db_path)
        with db:
            db.execute('UPDATE rows SET row_section=NULL WHERE path=?', (str(source.resolve()),))
        self.assertTrue(needs_search_update(db))
        db.close()
        self.assertEqual(sync(self.db_path, [str(self.root)], threading.Event())['indexed'], 1)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
        mixed = search(self.db_path, '"Company"')[1]
        self.assertEqual(len(mixed), 3)
        self.assertTrue(all(row['mixed_headers'] for row in mixed))
        # Generic headings must remain consistent even on a one-row page.
        self.assertTrue(search(self.db_path, '"Company"', page_size=1)[1][0]['mixed_headers'])
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtWidgets import QApplication, QCheckBox
        from local_register_search.ui import Window
        app = QApplication.instance() or QApplication([])
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        window.search_done(window.generation, (3, mixed))
        self.assertEqual(window.table.horizontalHeaderItem(3).text(), 'Column A')
        self.assertEqual(window.table.horizontalHeaderItem(window.table.columnCount() - 2).text(), 'Section')
        self.assertEqual(window.table.horizontalHeaderItem(window.table.columnCount() - 1).text(), 'Open')
        client_index = next(i for i, row in enumerate(mixed) if row['row_section'] == 'B - Client Data')
        self.assertEqual(window.row_fields(mixed[client_index])[0][1], 'Field')
        window.select_record(window.table.model().index(client_index, 0))
        self.assertIn('Section:', window.detail.toPlainText())
        self.assertIn('B - Client Data', window.detail.toPlainText())
        option = next(control for control in window.findChildren(QCheckBox) if control.text() == 'Include section heading in search')
        self.assertTrue(option.isChecked())
        option.setChecked(False)
        window.timer.stop()
        self.assertFalse(window.include_section)
        records = search(self.db_path, '"Company" "B - Client Data"')[1]
        self.assertEqual(export_records(self.root / 'section-results.csv', records), 1)
        self.assertIn('B - Client Data', (self.root / 'section-results.csv').read_text(encoding='utf-8-sig'))

    def test_phrase_chips_colors_and_empty_cells(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        from PySide6.QtWidgets import QApplication
        from local_register_search.ui import Window
        from local_register_search.chip_input import quoted
        app = QApplication.instance() or QApplication([])
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        bar = window.query
        phrases = ['Administration - Bank Guarantees', 'CATEGORIES (Structured)']
        for phrase in phrases:
            bar.editor.setText(phrase)
            QTest.keyClick(bar.editor, Qt.Key.Key_Return)
        self.assertEqual([term for term, color in bar.chips], phrases)
        self.assertEqual(parse_query(bar.text()), [phrase.casefold() for phrase in phrases])
        self.assertEqual(bar.editor.placeholderText(), 'Add another keyword...')
        row = search(self.db_path, 'rishikesh')[1][0]
        row['cells'] = json.dumps({'A': phrases[0], 'B': phrases[1], 'C': ''})
        row['headers'] = json.dumps({'A': 'Department', 'B': 'Category', 'C': 'Empty'})
        db = connect(self.db_path)
        from local_register_search.search_terms import index_text
        text, digits = index_text(json.loads(row['cells']))
        with db:
            db.execute('UPDATE rows SET cells=?,row_text=?,row_digits=? WHERE id=?', (row['cells'], text, digits, row['id']))
        db.close()
        self.assertEqual(search(self.db_path, bar.text())[0], 1)
        window.search_done(window.generation, (1, [row]))
        self.assertIn(bar.chips[0][1], window.highlight(phrases[0]))
        self.assertIn(bar.chips[1][1], window.highlight(phrases[1]))
        self.assertEqual(window.highlight('-'), '-')
        self.assertEqual(window.table.item(0, 5).data(Qt.ItemDataRole.UserRole), 'empty-cell')
        self.assertIn(phrases[0], window.result_count.text())
        bar.edit_chip(0)
        bar.editor.setText('Administration - Bank Guarantees updated')
        bar.enter()
        self.assertEqual(bar.chips[0][0], 'Administration - Bank Guarantees updated')
        bar.load_query(' '.join(quoted(term) for term in phrases))
        self.assertEqual([term for term, _ in bar.chips], phrases)
        # Avoid starting workers while testing removal mechanics.
        bar.searchRequested.disconnect(window.submit_search)
        QTest.keyClick(bar.editor, Qt.Key.Key_Backspace)
        self.assertEqual(len(bar.chips), 1)
        bar.clear()
        self.assertEqual(window.table.rowCount(), 0)
        bar.editor.setText('-')
        bar.enter()
        self.assertEqual(bar.chips, [])
        for i in range(9):
            bar.editor.setText('keyword ' + str(i))
            bar.enter()
        self.assertEqual(len(bar.chips), 8)
        bar.clear()
        bar.editor.setText('a')
        self.assertEqual(bar.text(), '')
        bar.enter()
        self.assertEqual(parse_query(bar.text()), ['a'])

    def test_referenced_file_links_only_for_file_cells(self):
        from PySide6.QtCore import QUrl
        from PySide6.QtWidgets import QApplication
        from local_register_search.ui import Window
        from local_register_search.desktop_features import referenced_file
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        app = QApplication.instance() or QApplication([])
        pdf = self.root / 'drawing.pdf'
        pdf.write_bytes(b'%PDF-1.4')
        self.assertIsNone(referenced_file('N/A', self.root / 'register.xlsx'))
        self.assertIsNone(referenced_file('Company Pvt. Ltd', self.root / 'register.xlsx'))
        self.assertEqual(referenced_file('drawing.pdf', self.root / 'register.xlsx'), pdf)
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        row = search(self.db_path, 'rishikesh')[1][0]
        row['cells'] = json.dumps({'A': 'N/A', 'B': str(pdf), 'C': 'Company Pvt. Ltd'})
        window.search_done(window.generation, (1, [row]))
        self.assertNotIn('Open file', window.detail.toPlainText())
        self.assertFalse(window.open_file_button.isHidden())
        window.select_record(window.table.model().index(0, 0))
        window.show()
        app.processEvents()
        self.assertEqual({button.height() for button in window.record_buttons}, {28})
        primary, *secondary = window.record_buttons
        self.assertLessEqual(max(button.width() for button in secondary) - min(button.width() for button in secondary), 1)
        self.assertGreater(primary.width(), secondary[0].width())
        self.assertLess(primary.y(), window.open_file_button.y())
        self.assertLess(window.open_file_button.y(), window.detail.y())
        self.assertEqual(window.detail_file_paths, {'B': pdf})
        with patch('local_register_search.ui.QDesktopServices.openUrl', return_value=True) as opened:
            window.open_file_button.click()
            opened.assert_called_once_with(QUrl.fromLocalFile(str(pdf)))
        with patch('local_register_search.ui.QDesktopServices.openUrl') as opened:
            window.open_referenced_file(QUrl('record-file:A'))
            opened.assert_not_called()
        row['cells'] = json.dumps({'A': 'N/A'})
        window.search_done(window.generation, (1, [row]))
        self.assertTrue(window.open_file_button.isHidden())
        row['cells'] = json.dumps({'A': 'needle', 'B': 'NOTE: ' + 'long note text ' * 30})
        row['headers'] = json.dumps({'A': 'Name', 'B': 'Note'})
        window.search_done(window.generation, (1, [row]))
        window.field_filter.setText('needle')
        self.assertIn('needle', window.detail.toPlainText())
        self.assertNotIn('long note text', window.detail.toPlainText())
        self.assertEqual(window.field_count.text(), '1 of 2 fields')
        window.field_filter.clear()
        self.assertIn('Show more', window.detail.toPlainText())
        window.detail_link_clicked(QUrl('expand-note:B'))
        self.assertIn('Show less', window.detail.toPlainText())
        self.assertIn('long note text ' * 30, window.detail.toPlainText())
        header_note = 'NOTE: If a hyperlink is not opening, please change the default application from Chrome to a specific PDF viewer. ' * 4
        row['headers'] = json.dumps({'A': header_note})
        row['cells'] = json.dumps({'A': '3999'})
        window.search_done(window.generation, (1, [row]))
        rendered = window.detail.toHtml()
        self.assertIn('colspan="2"', rendered)
        self.assertIn('Column A', window.detail.toPlainText())
        self.assertIn('3999', window.detail.toPlainText())
        self.assertIn('Show more', window.detail.toPlainText())
        window.detail_link_clicked(QUrl('expand-note:A:header'))
        self.assertIn(header_note.strip(), window.detail.toPlainText())

    def test_table_readability_and_saved_widths(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QFontDatabase, QFontMetrics
        from PySide6.QtWidgets import QApplication
        from local_register_search.ui import Window
        app = QApplication.instance() or QApplication([])
        font_file = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'segoeui.ttf'
        if font_file.exists():
            QFontDatabase.addApplicationFont(str(font_file))
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        window.show()
        long_text = 'Learning and Understanding ' * 15
        row = search(self.db_path, 'rishikesh')[1][0]
        row['cells'] = json.dumps({'A': long_text, 'B': '2973.6', 'C': '', 'D': 'short'})
        row['headers'] = json.dumps({'A': 'KPI', 'B': 'Weightage', 'C': '', 'D': 'Target', 'E': 'Empty'})
        row['path'] = str(self.root / 'PMS draft KPI Sheet with a very long file name.xlsx')
        window.search_done(window.generation, (1, [row]))
        app.processEvents()
        self.assertEqual(window.table.horizontalHeaderItem(3).text(), 'KPI')
        self.assertEqual(window.table.horizontalHeaderItem(5).text(), 'Column C')
        self.assertEqual(window.table.item(0, 5).text(), '-')
        self.assertEqual(window.table.item(0, 7).text(), '-')
        self.assertTrue(window.table.item(0, 4).textAlignment() & Qt.AlignmentFlag.AlignRight)
        self.assertTrue(window.table.item(0, 3).textAlignment() & Qt.AlignmentFlag.AlignLeft)
        font = window.table.font()
        lines = window.match_delegate.visible_lines(long_text, font, 110, 3)
        self.assertEqual(len(lines), 1)
        self.assertTrue(lines[-1].endswith('…'))
        self.assertGreaterEqual(window.table.rowHeight(0), 18)
        self.assertLessEqual(window.table.rowHeight(0), 22)
        self.assertEqual(window.table.rowHeight(0), window.table.frozen.rowHeight(0))
        self.assertEqual(window.table.item(0, 3).toolTip(), long_text)
        window.table.setColumnWidth(0, 400)
        self.assertLessEqual(window.table.columnWidth(0), 260)
        open_button = window.table.cellWidget(0, window.table.columnCount() - 1)
        with patch.object(window, 'open_path') as open_path:
            open_button.click()
            open_path.assert_called_once_with(False, row)
        existing_row = dict(row, path=str(self.root / 'register.xlsx'))
        with patch('local_register_search.desktop_features.open_excel_record', return_value=False) as jump, \
             patch('local_register_search.ui.QDesktopServices.openUrl', return_value=True) as fallback:
            window.open_path(False, existing_row)
            jump.assert_called_once_with(Path(existing_row['path']), existing_row['sheet'], existing_row['rownum'])
            fallback.assert_called_once()
        filename = Path(row['path']).name
        elided = window.match_delegate.visible_lines(filename, font, 120, 0)[0]
        self.assertNotEqual(elided, filename)
        self.assertTrue(elided.endswith('.xlsx'))
        window.select_record(window.table.model().index(0, 0))
        self.assertIn(long_text, window.detail.toPlainText())
        window.table.setColumnWidth(3, 330)
        window.column_resize_timer.stop()
        window.finish_column_resize()
        window.search_done(window.generation, (1, [row]))
        self.assertEqual(window.table.columnWidth(3), 330)
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            reopened = Window()
        self.addCleanup(reopened.close)
        reopened.search_done(reopened.generation, (1, [row]))
        self.assertEqual(reopened.table.columnWidth(3), 330)

    def test_dark_theme_page_backgrounds_and_button_colors(self):
        from PySide6.QtWidgets import QApplication, QLabel
        from local_register_search.ui import Window
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        app = QApplication.instance() or QApplication([])
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        window.show()
        for dark, expected in [(True, '#111e2a'), (False, '#f8fafb'), (True, '#111e2a')]:
            window.set_dark_mode(dark)
            for index in (1, 2, 3):
                window.show_tab(index)
                app.processEvents()
                content = window.pages.widget(index).widget()
                image = content.grab().toImage()
                self.assertEqual(image.pixelColor(image.width() - 3, image.height() - 3).name(), expected)
            online = next(label for label in window.findChildren(QLabel) if 'Offline' in label.text())
            self.assertIn('#74db99', online.styleSheet())
        window.clear_all_button.setEnabled(True)
        app.processEvents()
        image = window.clear_all_button.grab().toImage()
        self.assertEqual(image.pixelColor(5, image.height() // 2).name(), '#462d35')
        danger = next(label for label in window.findChildren(QLabel) if label.text() == 'Danger zone')
        self.assertIn('#ffb4ab', danger.styleSheet())

    def test_top_navigation_drawer_responsive_and_confirmations(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QFontDatabase
        from PySide6.QtWidgets import QApplication, QMessageBox
        from PySide6.QtTest import QTest
        from local_register_search.ui import Window
        app = QApplication.instance() or QApplication([])
        font_file = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'segoeui.ttf'
        if font_file.exists():
            QFontDatabase.addApplicationFont(str(font_file))
        with patch('local_register_search.ui.data_dir', return_value=self.root):
            window = Window()
        self.addCleanup(window.close)
        window.show()
        app.processEvents()
        self.assertIsNone(window.findChild(type(window.centralWidget()), 'sidebar'))
        self.assertEqual(window.pages.count(), 4)
        self.assertTrue(window.detail_drawer.isHidden())
        self.assertEqual(window.table.horizontalScrollBarPolicy(), Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        window.folders = [str(self.root)]
        window.save_folders()
        window.folder_filter.setCurrentIndex(1)
        self.assertEqual(window.folder_list.currentRow(), 0)
        window.nav_buttons[1].click()
        self.assertEqual(window.pages.currentIndex(), 1)
        self.assertTrue(window.nav_buttons[1].isChecked())
        self.assertFalse(window.nav_buttons[0].isChecked())
        window.show_tab(0)
        window.query.setText('rishikesh')
        window.timer.stop()
        window.search_done(window.generation, search(self.db_path, 'rishikesh'))
        self.assertTrue(window.detail_drawer.isHidden())
        window.select_record(window.table.model().index(0, 0))
        app.processEvents()
        self.assertFalse(window.detail_drawer.isHidden())
        self.assertIn('Name', window.detail.toPlainText())
        window.drawer_close.click()
        self.assertTrue(window.detail_drawer.isHidden())
        window.resize(640, 600)
        app.processEvents()
        top = window.findChild(type(window.centralWidget()), 'topbar')
        self.assertGreaterEqual(top.height(), 40)
        for control in window.nav_buttons:
            self.assertLessEqual(control.mapTo(top, control.rect().bottomRight()).x(), top.width())
        self.assertLessEqual(window.query.geometry().right(), window.query.parentWidget().width())
        with patch('local_register_search.ui.QMessageBox.question', return_value=QMessageBox.StandardButton.No) as confirm:
            window.delete_folder_data()
            window.delete_index()
            self.assertEqual(confirm.call_count, 2)
        self.assertEqual(search(self.db_path, 'rishikesh')[0], 1)
        window.scanning = True
        window.update_progress(120, 500, 'register.xlsx')
        QTest.qWait(180)
        self.assertFalse(window.sync_panel.isHidden())
        self.assertEqual(window.progress_label.text(), '120/500 files')
        window.scanning = False
        window.resize(1440, 900)
        QTest.qWait(700)
        app.processEvents()
        if os.environ.get('SEARCH_UI_PREVIEW'):
            window.grab().save(os.environ['SEARCH_UI_PREVIEW'])


if __name__ == '__main__':
    unittest.main()
