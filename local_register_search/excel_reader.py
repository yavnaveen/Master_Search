import datetime
import re
from zipfile import ZipFile
from xml.etree.ElementTree import iterparse
from pathlib import Path
import openpyxl
import xlrd
from openpyxl.utils import get_column_letter


class Cancelled(Exception):
    pass


def display(value, fmt=''):
    if value is None:
        return ''
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat(sep=' ') if isinstance(value, datetime.datetime) else value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if re.fullmatch('0+', fmt) and float(value).is_integer():
            return str(int(value)).zfill(len(fmt))
    return str(value)


def merged_rows(path, worksheet_path, cancel):
    rows = set()
    from openpyxl.utils.cell import range_boundaries
    with ZipFile(path) as archive, archive.open(worksheet_path.lstrip('/')) as stream:
        for event, element in iterparse(stream, events=('end',)):
            if cancel.is_set():
                raise Cancelled()
            if element.tag.endswith('}mergeCell'):
                left, top, right, bottom = range_boundaries(element.attrib['ref'])
                if right > left:
                    rows.add(top)
            element.clear()
    return rows


def read_rows(path, cancel, with_context=False):
    """Yield actual row references; never evaluate formulas or save a workbook."""
    if Path(path).suffix.lower() == '.xls':
        book = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
        try:
            for sheet_name in book.sheet_names():
                sheet = book.sheet_by_name(sheet_name)
                for r in range(sheet.nrows):
                    if cancel.is_set():
                        raise Cancelled()
                    cells = {}
                    styled = False
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        fmt = book.format_map.get(book.xf_list[cell.xf_index].format_key)
                        text = display(value, fmt.format_str if fmt else '')
                        if text:
                            cells[get_column_letter(c + 1)] = text
                            xf = book.xf_list[cell.xf_index]
                            styled |= bool(book.font_list[xf.font_index].bold or xf.background.fill_pattern)
                    if cells:
                        merged = any(top == r and right - left > 1 for top, bottom, left, right in sheet.merged_cells) if with_context else False
                        record = (sheet.name, r + 1, bool(sheet.visibility), cells)
                        yield record + ({'styled': styled, 'merged': merged},) if with_context else record
                book.unload_sheet(sheet.name)
        finally:
            book.release_resources()
        return
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = None
    try:
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        for sheet in cached:
            raw = formulas[sheet.title]
            merges = merged_rows(path, raw._worksheet_path, cancel) if with_context else set()
            sheet.reset_dimensions()
            raw.reset_dimensions()
            for rownum, (values, originals) in enumerate(zip(sheet.iter_rows(), raw.iter_rows()), 1):
                if cancel.is_set():
                    raise Cancelled()
                cells = {}
                styled = False
                for col, (cell, original) in enumerate(zip(values, originals), 1):
                    text = display(cell.value, getattr(original, 'number_format', ''))
                    if original.data_type == 'f' and cell.value is None:
                        text = '[Formula: cached value unavailable] ' + str(original.value)
                    if text:
                        cells[get_column_letter(col)] = text
                        styled |= bool(original.font.bold or original.fill.patternType)
                if cells:
                    record = (sheet.title, rownum, sheet.sheet_state != 'visible', cells)
                    yield record + ({'styled': styled, 'merged': rownum in merges},) if with_context else record
    finally:
        cached.close()
        if formulas:
            formulas.close()
